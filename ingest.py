"""
ingest.py
---------
Run this ONCE (and again whenever your source documents change) to build
the FAISS vector index and save it to disk. main.py then just loads the
saved index at startup instead of rebuilding it on every launch.

Usage:
    source venv/bin/activate
    python ingest.py
"""

import os
from pathlib import Path

import pandas as pd
from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings

# ---------------------------------------------------------------------
# CONFIG - edit these for your server
# ---------------------------------------------------------------------
ROOT_FOLDER = os.environ.get("RAG_DOCS_FOLDER", "./docs")
INDEX_FOLDER = os.environ.get("RAG_INDEX_FOLDER", "./faiss_index")
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
# ---------------------------------------------------------------------


def load_documents(root_folder: str) -> list[Document]:
    documents = []
    root = Path(root_folder)

    if not root.exists():
        raise Exception(f"Folder '{root_folder}' not found.")

    print("Scanning files...")

    for file in root.rglob("*"):
        if not file.is_file():
            continue

        try:
            text = ""
            suffix = file.suffix.lower()

            if suffix in (".txt", ".md"):
                text = file.read_text(encoding="utf-8", errors="ignore")

            elif suffix == ".pdf":
                pdf = PdfReader(str(file))
                pages = [page.extract_text() or "" for page in pdf.pages]
                text = "\n".join(pages)

            elif suffix == ".docx":
                doc = DocxDocument(str(file))
                text = "\n".join(p.text for p in doc.paragraphs)

            elif suffix == ".csv":
                df = pd.read_csv(file, dtype=str, encoding_errors="ignore")
                text = df.to_string()

            else:
                continue

            text = text.strip()
            if not text:
                continue

            documents.append(
                Document(page_content=text, metadata={"source": str(file)})
            )

        except Exception as e:
            print(f"Skipped {file}: {e}")

    print(f"Loaded {len(documents)} documents")
    return documents


def main():
    documents = load_documents(ROOT_FOLDER)

    if len(documents) == 0:
        raise Exception("No readable documents found.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    chunks = splitter.split_documents(documents)
    print(f"Created {len(chunks)} chunks")

    if len(chunks) == 0:
        raise Exception("No chunks generated.")

    print("Loading embedding model...")
    embeddings = HuggingFaceEmbeddings(model_name=EMBEDDING_MODEL)

    print("Building FAISS index (this can take a while)...")
    vectorstore = FAISS.from_documents(chunks, embeddings)

    Path(INDEX_FOLDER).mkdir(parents=True, exist_ok=True)
    vectorstore.save_local(INDEX_FOLDER)

    print(f"Index saved to: {INDEX_FOLDER}")
    print("Done. You can now start the server (main.py / systemd service).")


if __name__ == "__main__":
    main()
